from django.shortcuts import (
    render,
    redirect,
    get_object_or_404
)

from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.contrib import messages
from django.http import JsonResponse
from django.db import transaction
from django.utils import timezone

from openpyxl import load_workbook

from .models import UserProfile
from .forms import (
    UserCreateForm,
    UserEditForm
)

from calling.models import CallingTask


# =========================================================
# USER MANAGEMENT PERMISSION
# =========================================================

def can_manage_users(user):

    if not user.is_authenticated:
        return False

    if user.is_superuser:
        return True

    try:
        return user.profile.role == 'Admin'

    except UserProfile.DoesNotExist:
        return False


# =========================================================
# USER LIST
# =========================================================

@login_required
def user_list(request):

    if not can_manage_users(request.user):
        return redirect('dashboard')

    users = User.objects.select_related(
        'profile'
    ).order_by(
        'first_name',
        'username'
    )

    form = UserCreateForm()

    context = {
        'users': users,
        'form': form,
    }

    return render(
        request,
        'accounts/user_list.html',
        context
    )


# =========================================================
# ADD USER
# =========================================================

@login_required
def add_user(request):

    if not can_manage_users(request.user):
        return redirect('dashboard')

    if request.method != 'POST':
        return redirect('user_list')

    form = UserCreateForm(
        request.POST
    )

    if form.is_valid():

        full_name = (
            form.cleaned_data[
                'full_name'
            ]
            .strip()
            .split(' ', 1)
        )

        first_name = full_name[0]

        last_name = (
            full_name[1]
            if len(full_name) > 1
            else ''
        )

        user = User.objects.create_user(
            username=form.cleaned_data[
                'username'
            ],
            email=form.cleaned_data[
                'email'
            ],
            password=form.cleaned_data[
                'password'
            ],
            first_name=first_name,
            last_name=last_name,
        )

        UserProfile.objects.create(
            user=user,
            role=form.cleaned_data[
                'role'
            ],
            contact_no=form.cleaned_data[
                'contact_no'
            ],
        )

        messages.success(
            request,
            'User created successfully.'
        )

        return redirect(
            'user_list'
        )

    users = User.objects.select_related(
        'profile'
    ).order_by(
        'first_name',
        'username'
    )

    return render(
        request,
        'accounts/user_list.html',
        {
            'users': users,
            'form': form,
            'open_add_modal': True,
        }
    )


# =========================================================
# EDIT USER
# =========================================================

@login_required
def edit_user(request, pk):

    if not can_manage_users(request.user):
        return redirect('dashboard')

    user = get_object_or_404(
        User,
        pk=pk
    )

    # Super Admin cannot be edited here.
    if user.is_superuser:
        messages.error(
            request,
            'Super Admin cannot be edited.'
        )

        return redirect(
            'user_list'
        )

    profile, created = (
        UserProfile.objects.get_or_create(
            user=user,
            defaults={
                'role': 'Counsellor',
                'contact_no': '',
            }
        )
    )

    if request.method == 'POST':

        form = UserEditForm(
            request.POST,
            user_instance=user
        )

        if form.is_valid():

            full_name = (
                form.cleaned_data[
                    'full_name'
                ]
                .strip()
                .split(' ', 1)
            )

            user.first_name = (
                full_name[0]
            )

            user.last_name = (
                full_name[1]
                if len(full_name) > 1
                else ''
            )

            user.username = (
                form.cleaned_data[
                    'username'
                ]
            )

            user.email = (
                form.cleaned_data[
                    'email'
                ]
            )

            user.save()

            profile.contact_no = (
                form.cleaned_data[
                    'contact_no'
                ]
            )

            profile.role = (
                form.cleaned_data[
                    'role'
                ]
            )

            profile.save()

            messages.success(
                request,
                'User updated successfully.'
            )

        else:

            messages.error(
                request,
                'Unable to update user. '
                'Please check the entered details.'
            )

    return redirect(
        'user_list'
    )


# =========================================================
# RESET PASSWORD
# =========================================================

@login_required
def reset_password(request, pk):

    if not can_manage_users(request.user):
        return redirect('dashboard')

    user = get_object_or_404(
        User,
        pk=pk
    )

    if user.is_superuser:

        messages.error(
            request,
            'Super Admin password cannot be reset here.'
        )

        return redirect(
            'user_list'
        )

    if request.method == 'POST':

        password = request.POST.get(
            'password',
            ''
        )

        confirm_password = request.POST.get(
            'confirm_password',
            ''
        )

        if not password:

            messages.error(
                request,
                'Password is required.'
            )

        elif password != confirm_password:

            messages.error(
                request,
                'Passwords do not match.'
            )

        else:

            user.set_password(
                password
            )

            user.save()

            messages.success(
                request,
                'Password reset successfully.'
            )

    return redirect(
        'user_list'
    )


# =========================================================
# ACTIVATE / DEACTIVATE USER
# =========================================================

@login_required
def toggle_user_status(request, pk):

    if not can_manage_users(request.user):
        return redirect('dashboard')

    user = get_object_or_404(
        User,
        pk=pk
    )

    if user.is_superuser:

        messages.error(
            request,
            'Super Admin status cannot be changed.'
        )

        return redirect(
            'user_list'
        )

    if request.method == 'POST':

        user.is_active = (
            not user.is_active
        )

        user.save(
            update_fields=[
                'is_active'
            ]
        )

        if user.is_active:

            messages.success(
                request,
                'User activated successfully.'
            )

        else:

            messages.success(
                request,
                'User deactivated successfully.'
            )

    return redirect(
        'user_list'
    )


@login_required
def delete_user(request, pk):
    if not can_manage_users(request.user):
        return redirect('dashboard')
    if request.method != 'POST':
        return redirect('user_list')

    user = get_object_or_404(User, pk=pk)
    if user.is_superuser:
        messages.error(request, 'Super Admin accounts cannot be deleted here.')
    elif user.pk == request.user.pk:
        messages.error(request, 'You cannot delete your own account.')
    else:
        username = user.get_full_name() or user.username
        user.delete()
        messages.success(
            request,
            f'{username} was deleted. Candidate and calling-task records were retained; their calls are now under Unassigned / Archived.'
        )
    return redirect('user_list')


# =========================================================
# ASSIGN CALLING TASKS
# =========================================================

@login_required
def assign_tasks(request):
    if not can_manage_users(request.user):
        return redirect('dashboard')

    tab = request.GET.get('tab', 'counsellor')
    if tab not in {'counsellor', 'trainer'}:
        tab = 'counsellor'

    counsellors = User.objects.filter(
        is_active=True, is_superuser=False, profile__role='Counsellor'
    ).select_related('profile').order_by('first_name', 'username')
    trainers = User.objects.filter(
        is_active=True, is_superuser=False, profile__role='Trainer'
    ).select_related('profile').order_by('first_name', 'username')

    all_tasks = CallingTask.objects.all()
    selected_counsellor = None
    selected_key = request.GET.get('counsellor', '')
    tasks = CallingTask.objects.none()
    if tab == 'counsellor' and selected_key == 'unassigned':
        tasks = all_tasks.filter(counsellor__isnull=True)
    elif tab == 'counsellor' and selected_key.isdigit():
        selected_counsellor = counsellors.filter(pk=int(selected_key)).first()
        if selected_counsellor:
            tasks = all_tasks.filter(counsellor=selected_counsellor)

    context = {
        'tab': tab,
        'counsellors': counsellors,
        'trainers': trainers,
        'selected_counsellor': selected_counsellor,
        'selected_key': selected_key,
        'tasks': tasks.select_related('candidate').order_by('-assigned_date', '-id'),
        'total_calls': all_tasks.count(),
        'pending_calls': all_tasks.filter(status='Pending').count(),
        'completed_calls': all_tasks.filter(status='Completed').count(),
        'revoked_calls': all_tasks.filter(status='Revoked').count(),
        'archived_calls': all_tasks.filter(counsellor__isnull=True).count(),
        'transfer_targets': counsellors,
    }
    return render(request, 'accounts/assign_tasks.html', context)


@login_required
def edit_call_task(request, pk):
    if not can_manage_users(request.user):
        return redirect('dashboard')
    task = get_object_or_404(CallingTask, pk=pk)
    if request.method == 'POST':
        name = request.POST.get('candidate_name', '').strip()
        mobile = ''.join(char for char in request.POST.get('mobile', '') if char.isdigit())
        if task.status != 'Pending':
            messages.error(request, 'Only pending calls can be modified.')
        elif not name or len(mobile) != 10:
            messages.error(request, 'Enter a candidate name and a valid 10-digit mobile number.')
        elif CallingTask.objects.exclude(pk=task.pk).exclude(status='Revoked').filter(mobile=mobile).exists():
            messages.error(request, 'This number is already assigned in another calling task.')
        else:
            task.candidate_name = name
            task.mobile = mobile
            task.reg_no = request.POST.get('reg_no', '').strip()
            task.reference = request.POST.get('reference', '').strip()
            task.save()
            messages.success(request, 'Calling task updated.')
        return redirect(f"/accounts/assign-tasks/?tab=counsellor&counsellor={task.counsellor_id or 'unassigned'}")
    return render(request, 'accounts/edit_call_task.html', {'task': task})


@login_required
def transfer_call_task(request, pk):
    if not can_manage_users(request.user):
        return redirect('dashboard')
    task = get_object_or_404(CallingTask, pk=pk)
    selected_key = task.counsellor_id or 'unassigned'
    if request.method == 'POST':
        target = get_object_or_404(
            User.objects.select_related('profile'),
            pk=request.POST.get('counsellor_id'), is_active=True,
            is_superuser=False, profile__role='Counsellor'
        )
        if task.status != 'Pending':
            messages.error(request, 'Only pending calls can be transferred.')
        elif task.counsellor_id == target.pk:
            messages.info(request, 'This call is already assigned to that counsellor.')
        else:
            task.counsellor = target
            task.assigned_date = timezone.localdate()
            task.save()
            if task.candidate_id:
                task.candidate.counsellor = target
                task.candidate.save(update_fields=['counsellor', 'updated_at'])
            messages.success(request, f'{task.candidate_name} was transferred to {target.get_full_name() or target.username}.')
        return redirect(f'/accounts/assign-tasks/?tab=counsellor&counsellor={selected_key}')
    return redirect('assign_tasks')


@login_required
def revoke_call_task(request, pk):
    if not can_manage_users(request.user):
        return redirect('dashboard')
    task = get_object_or_404(CallingTask, pk=pk)
    selected_key = task.counsellor_id or 'unassigned'
    if request.method == 'POST':
        if task.status == 'Pending':
            task.status = 'Revoked'
            task.save(update_fields=['status', 'updated_at'])
            messages.success(request, 'Call task revoked. Its record has been retained.')
        else:
            messages.error(request, 'Only pending calls can be revoked.')
    return redirect(f'/accounts/assign-tasks/?tab=counsellor&counsellor={selected_key}')

@login_required
def preview_call_upload(request, pk):
    if not can_manage_users(request.user):
        return JsonResponse({'error': 'Permission denied.'}, status=403)
    if request.method != 'POST':
        return JsonResponse({'error': 'Use POST to preview a file.'}, status=405)

    counsellor = get_object_or_404(User, pk=pk, is_active=True)
    try:
        if counsellor.profile.role != 'Counsellor':
            return JsonResponse({'error': 'Calls can only be assigned to Counsellors.'}, status=400)
    except UserProfile.DoesNotExist:
        return JsonResponse({'error': 'This user does not have a valid profile.'}, status=400)

    excel_file = request.FILES.get('excel_file')
    if not excel_file or not excel_file.name.lower().endswith('.xlsx'):
        return JsonResponse({'error': 'Choose a valid .xlsx Excel file.'}, status=400)

    try:
        workbook = load_workbook(excel_file, read_only=True, data_only=True)
        worksheet = workbook.active
        first_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not first_row:
            return JsonResponse({'error': 'The Excel file is empty.'}, status=400)

        normalize = lambda value: ''.join(str(value or '').strip().lower().replace('.', '').split()).replace('_', '')
        headers = [normalize(value) for value in first_row]
        aliases = {
            'name': {'name', 'candidatename'},
            'mobile': {'mobile', 'mobileno', 'mobilenumber', 'contact', 'contactno', 'contactnumber'},
        }
        name_col = next((i for i, value in enumerate(headers) if value in aliases['name']), None)
        mobile_col = next((i for i, value in enumerate(headers) if value in aliases['mobile']), None)
        if name_col is None or mobile_col is None:
            return JsonResponse({'error': 'Required columns Name and Mobile were not found. Use the sample Excel format.'}, status=400)

        def clean(value):
            if value is None:
                return ''
            if isinstance(value, float) and value.is_integer():
                return str(int(value))
            return str(value).strip()

        counts = {'rows': 0, 'valid': 0, 'invalid': 0, 'duplicates': 0, 'blank': 0}
        examples = []
        seen_mobiles = set()
        existing_numbers = set(CallingTask.objects.exclude(status='Revoked').values_list('mobile', flat=True))
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            counts['rows'] += 1
            name = clean(row[name_col]) if name_col < len(row) else ''
            raw_mobile = clean(row[mobile_col]) if mobile_col < len(row) else ''
            if not name and not raw_mobile:
                counts['blank'] += 1
                continue
            mobile = ''.join(char for char in raw_mobile if char.isdigit())
            if not name or len(mobile) != 10:
                counts['invalid'] += 1
                continue
            if mobile in existing_numbers or mobile in seen_mobiles:
                counts['duplicates'] += 1
                continue
            seen_mobiles.add(mobile)
            counts['valid'] += 1
            if len(examples) < 5:
                examples.append({'name': name, 'mobile': mobile})

        return JsonResponse({'counts': counts, 'examples': examples})
    except Exception:
        return JsonResponse({'error': 'Unable to read this Excel file. Please check the file and try again.'}, status=400)
    finally:
        if 'workbook' in locals():
            workbook.close()

@login_required
def assign_calls(request, pk):

    if not can_manage_users(request.user):
        return redirect('dashboard')

    counsellor = get_object_or_404(
        User.objects.select_related(
            'profile'
        ),
        pk=pk,
        is_active=True
    )

    # Only Counsellor users can receive calling tasks.
    if counsellor.is_superuser:

        messages.error(
            request,
            'Calling tasks can only be assigned '
            'to Counsellors.'
        )

        return redirect(
            f'/accounts/assign-tasks/?tab=counsellor&counsellor={pk}'
        )

    try:

        if counsellor.profile.role != 'Counsellor':

            messages.error(
                request,
                'Calling tasks can only be assigned '
                'to Counsellors.'
            )

            return redirect(
                f'/accounts/assign-tasks/?tab=counsellor&counsellor={pk}'
            )

    except UserProfile.DoesNotExist:

        messages.error(
            request,
            'This user does not have a valid profile.'
        )

        return redirect(
            f'/accounts/assign-tasks/?tab=counsellor&counsellor={pk}'
        )

    if request.method != 'POST':

        return redirect(
            f'/accounts/assign-tasks/?tab=counsellor&counsellor={pk}'
        )

    excel_file = request.FILES.get(
        'excel_file'
    )

    if not excel_file:

        messages.error(
            request,
            'Please select an Excel file.'
        )

        return redirect(
            f'/accounts/assign-tasks/?tab=counsellor&counsellor={pk}'
        )

    if not excel_file.name.lower().endswith(
        '.xlsx'
    ):

        messages.error(
            request,
            'Please upload a valid .xlsx Excel file.'
        )

        return redirect(
            f'/accounts/assign-tasks/?tab=counsellor&counsellor={pk}'
        )

    try:

        workbook = load_workbook(
            excel_file,
            read_only=True,
            data_only=True
        )

        worksheet = workbook.active

    except Exception:

        messages.error(
            request,
            'Unable to read the Excel file. '
            'Please upload a valid .xlsx file.'
        )

        return redirect(
            f'/accounts/assign-tasks/?tab=counsellor&counsellor={pk}'
        )


    # =====================================================
    # VALIDATE HEADER
    # =====================================================

    first_row = next(
        worksheet.iter_rows(
            min_row=1,
            max_row=1,
            values_only=True
        ),
        None
    )

    if not first_row:

        workbook.close()

        messages.error(
            request,
            'The Excel file is empty.'
        )

        return redirect(
            f'/accounts/assign-tasks/?tab=counsellor&counsellor={pk}'
        )


    def normalize_header(value):

        if value is None:
            return ''

        value = str(
            value
        ).strip().lower()

        value = value.replace(
            '.',
            ''
        )

        value = value.replace(
            ' ',
            ''
        )

        value = value.replace(
            '_',
            ''
        )

        return value


    headers = [
        normalize_header(value)
        for value in first_row
    ]


    # Supported:
    #
    # S.No | Reg. No | Ref | Name | Mobile

    header_aliases = {

        'sno': [
            'sno',
            'serialno',
            'serialnumber',
        ],

        'regno': [
            'regno',
            'registrationno',
            'registrationnumber',
        ],

        'ref': [
            'ref',
            'reference',
        ],

        'name': [
            'name',
            'candidatename',
        ],

        'mobile': [
            'mobile',
            'mobileno',
            'mobilenumber',
            'contact',
            'contactno',
            'contactnumber',
        ],
    }


    def find_column(field_name):

        aliases = header_aliases[
            field_name
        ]

        for index, header in enumerate(
            headers
        ):

            if header in aliases:
                return index

        return None


    reg_column = find_column(
        'regno'
    )

    ref_column = find_column(
        'ref'
    )

    name_column = find_column(
        'name'
    )

    mobile_column = find_column(
        'mobile'
    )


    # Name and Mobile are mandatory.
    if (
        name_column is None
        or
        mobile_column is None
    ):

        workbook.close()

        messages.error(
            request,
            'Invalid Excel format. '
            'Required columns are Name and Mobile. '
            'Expected format: '
            'S.No | Reg. No | Ref | Name | Mobile'
        )

        return redirect(
            f'/accounts/assign-tasks/?tab=counsellor&counsellor={pk}'
        )


    imported_count = 0

    skipped_count = 0

    duplicate_count = 0

    rows_read = 0

    blank_rows = 0


    # =====================================================
    # HELPER FOR EXCEL VALUES
    # =====================================================

    def clean_value(value):

        if value is None:
            return ''

        # Prevent Excel numeric values such as
        # 12345.0 becoming "12345.0".
        if isinstance(
            value,
            float
        ) and value.is_integer():

            return str(
                int(value)
            )

        return str(
            value
        ).strip()


    # =====================================================
    # IMPORT
    # =====================================================

    try:

        with transaction.atomic():

            existing_numbers = set(
                CallingTask.objects.exclude(status='Revoked').values_list('mobile', flat=True)
            )
            upload_numbers = set()

            for row in worksheet.iter_rows(
                min_row=2,
                values_only=True
            ):

                rows_read += 1

                name = ''

                mobile = ''

                reg_no = ''

                reference = ''


                if name_column < len(row):

                    name = clean_value(
                        row[
                            name_column
                        ]
                    )


                if mobile_column < len(row):

                    mobile = clean_value(
                        row[
                            mobile_column
                        ]
                    )


                if (
                    reg_column is not None
                    and
                    reg_column < len(row)
                ):

                    reg_no = clean_value(
                        row[
                            reg_column
                        ]
                    )


                if (
                    ref_column is not None
                    and
                    ref_column < len(row)
                ):

                    reference = clean_value(
                        row[
                            ref_column
                        ]
                    )


                # Completely blank row.
                if (
                    not name
                    and
                    not mobile
                    and
                    not reg_no
                    and
                    not reference
                ):

                    blank_rows += 1

                    continue


                # Name is mandatory.
                if not name:

                    skipped_count += 1

                    continue


                # Mobile is mandatory and exactly 10 digits.
                mobile = ''.join(
                    character
                    for character in mobile
                    if character.isdigit()
                )


                if len(mobile) != 10:

                    skipped_count += 1

                    continue


                # A number can only be assigned once across the calling team.
                if mobile in existing_numbers or mobile in upload_numbers:

                    duplicate_count += 1

                    continue

                upload_numbers.add(mobile)


                CallingTask.objects.create(
                    counsellor=counsellor,
                    reg_no=reg_no,
                    reference=reference,
                    candidate_name=name,
                    mobile=mobile,
                    status='Pending'
                )


                imported_count += 1


    except Exception:

        workbook.close()

        messages.error(
            request,
            'The calling task import failed. '
            'No records from this upload were saved.'
        )

        return redirect(
            f'/accounts/assign-tasks/?tab=counsellor&counsellor={pk}'
        )


    workbook.close()


    # =====================================================
    # RESULT MESSAGE
    # =====================================================

    counsellor_name = (
        counsellor.get_full_name()
        or
        counsellor.username
    )


    if imported_count == 0:

        messages.warning(
            request,
            (
                f'No new calling tasks were assigned to '
                f'{counsellor_name}. '
                f'Rows checked: {rows_read}; assigned: 0; '
                f'skipped: {skipped_count}; duplicates: {duplicate_count}; '
                f'blank rows: {blank_rows}.'
            )
        )

    else:

        messages.success(
            request,
            (
                f'{imported_count} calling task(s) '
                f'assigned to {counsellor_name}. '
                f'Rows checked: {rows_read}; '
                f'skipped: {skipped_count}; duplicates: {duplicate_count}; '
                f'blank rows: {blank_rows}.'
            )
        )


    return redirect(
        f'/accounts/assign-tasks/?tab=counsellor&counsellor={pk}'
    )

